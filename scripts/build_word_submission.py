from __future__ import annotations

from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
TEX_PATH = ROOT / "reports" / "paper" / "main.tex"
BUILD_DIR = ROOT / "reports" / "paper" / "build"
REFERENCE_PATH = BUILD_DIR / "word_reference.docx"
DOCX_PATH = BUILD_DIR / "main.docx"


def set_east_asia_font(style_or_run, font_name: str, size_pt: float | None = None, bold: bool | None = None) -> None:
    font = style_or_run.font
    font.name = font_name
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    element = style_or_run.element
    rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:cs"), font_name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    fld_text = OxmlElement("w:t")
    fld_text.text = "1"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(fld_text)
    run._r.append(fld_end)


def configure_reference_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.6)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "宋体", 12, False)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    title_style = doc.styles["Title"]
    set_east_asia_font(title_style, "黑体", 16, False)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.line_spacing = 1.0
    title_style.paragraph_format.space_after = Pt(12)

    heading1 = doc.styles["Heading 1"]
    set_east_asia_font(heading1, "黑体", 14, False)
    heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading1.paragraph_format.line_spacing = 1.0
    heading1.paragraph_format.space_before = Pt(6)
    heading1.paragraph_format.space_after = Pt(6)

    for style_name in ["Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        set_east_asia_font(style, "黑体", 12, False)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(3)

    if "Abstract" not in doc.styles:
        abstract_style = doc.styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract_style = doc.styles["Abstract"]
    abstract_style.base_style = normal
    set_east_asia_font(abstract_style, "宋体", 12, False)
    abstract_style.paragraph_format.line_spacing = 1.0

    header = section.header
    for paragraph in header.paragraphs:
        if paragraph.text:
            paragraph.clear()

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.clear()
    add_page_number(paragraph)


def create_reference_docx() -> None:
    doc = Document()
    configure_reference_document(doc)
    doc.save(REFERENCE_PATH)


def post_process_docx() -> None:
    doc = Document(DOCX_PATH)

    for section in doc.sections:
        section.top_margin = Cm(2.6)
        section.bottom_margin = Cm(2.6)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.6)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)
        for paragraph in section.header.paragraphs:
            if paragraph.text:
                paragraph.clear()
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.clear()
        add_page_number(paragraph)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "宋体", 12, False)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    set_east_asia_font(doc.styles["Title"], "黑体", 16, False)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Title"].paragraph_format.line_spacing = 1.0

    set_east_asia_font(doc.styles["Heading 1"], "黑体", 14, False)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.line_spacing = 1.0

    for style_name in ["Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        set_east_asia_font(style, "黑体", 12, False)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.0

    page_break_inserted = False
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.style.name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if paragraph.style.name in {"Normal", "Body Text", "First Paragraph", "Compact"}:
            paragraph.style = doc.styles["Normal"]

        for run in paragraph.runs:
            set_east_asia_font(run, "宋体", 12)

        if paragraph.text.strip().startswith("关键词") and not page_break_inserted:
            paragraph.paragraph_format.line_spacing = 1.0
            if index + 1 < len(doc.paragraphs):
                paragraph.runs[-1].add_break(WD_BREAK.PAGE)
                page_break_inserted = True

    doc.save(DOCX_PATH)


def build_docx() -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    create_reference_docx()

    extra_args = [
        f"--reference-doc={REFERENCE_PATH}",
        f"--resource-path={ROOT}",
        "--standalone",
        "--wrap=none",
    ]
    pypandoc.convert_file(str(TEX_PATH), "docx", outputfile=str(DOCX_PATH), extra_args=extra_args)
    post_process_docx()


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)